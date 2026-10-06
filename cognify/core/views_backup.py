import random
import os

import subprocess
from pathlib import Path

from django.contrib.auth.forms import PasswordChangeForm
from django.db.models import Q

from django.contrib import messages
from django.contrib.auth import authenticate
from django.contrib.auth import login
from django.contrib.auth import logout

from django.contrib.auth.models import User
from django.contrib.auth.decorators import login_required


from django.core.mail import send_mail

from django.shortcuts import render
from django.shortcuts import redirect
from django.shortcuts import get_object_or_404,redirect

from pypdf import PdfReader
from docx import Document
from pptx import Presentation

from .forms import RegisterForm
from .forms import LoginForm
from .forms import NotebookForm

from .models import Notebook
from .models import Source
from .models import GeneratedResult
from .models import UserProfile

from openai import OpenAI,RateLimitError
from django.conf import settings

from google import genai
from google.genai import types

import threading

from django.db import close_old_connections
from django.http import JsonResponse

from django.contrib.auth import update_session_auth_hash
from django.core.exceptions import ValidationError
from django.contrib.auth.password_validation import validate_password

from django.db.models import FileField

from django.views.decorators.http import require_POST

from .models import AskSource, AskOutput
from .services.source_processor import (
    process_url,
    SourceProcessingError,
)
from django.utils import timezone


def gemini_generate(system_prompt, user_prompt):
    if not settings.GEMINI_API_KEY:
        raise RuntimeError("GEMINI_API_KEY is not configured.")

    client = genai.Client(api_key=settings.GEMINI_API_KEY)

    response = client.models.generate_content(
        model=settings.GEMINI_MODEL,
        contents=user_prompt,
        config=types.GenerateContentConfig(
            system_instruction=system_prompt,
            temperature=0.2,
        ),
    )

    if not response.text:
        raise RuntimeError("Gemini returned an empty response.")

    return response.text.strip()

@login_required
def ask_cognify(request):
    return render(
        request,
        "core/ask_cognify.html"
    )


@login_required
def process_ask_source(request):

    if request.method != "POST":

        return JsonResponse(
            {
                "success": False,
                "error": "POST request required."
            },
            status=405
        )

    url = request.POST.get(
        "url",
        ""
    ).strip()

    if not url:

        return JsonResponse(
            {
                "success": False,
                "error": "Please enter a URL."
            },
            status=400
        )

    try:

        result = process_url(
            url
        )

    except SourceProcessingError as error:

        return JsonResponse(
            {
                "success": False,
                "error": str(error)
            },
            status=400
        )

    except Exception as error:

        return JsonResponse(
            {
                "success": False,
                "error": (
                    "Unable to process this source."
                )
            },
            status=500
        )

    source = AskSource.objects.create(

        user=request.user,

        url=result.get(
            "source_url",
            url
        ),

        source_type=result.get(
            "source_type",
            "other"
        ),

        title=result.get(
            "title",
            ""
        ),

        author=result.get(
            "author",
            ""
        ),

        description=result.get(
            "description",
            ""
        ),

        thumbnail=result.get(
            "thumbnail",
            ""
        ),

        text=result.get(
            "text",
            ""
        ),

        transcript=result.get(
            "transcript",
            ""
        ),

        metadata=result.get(
            "metadata",
            {}
        ),

        chapters=result.get(
            "chapters",
            []
        ),

        access_status=result.get(
            "access_status",
            "partial"
        ),

        access_message=result.get(
            "access_message",
            ""
        ),

        retrieved_at=timezone.now()
    )

    return JsonResponse(
        {
            "success": True,

            "source_id": source.id,

            "title": source.title
            or "Untitled source",

            "source_type": source.get_source_type_display(),

            "url": source.url,

            "access_status": source.access_status,

            "access_message": (
                source.access_message
                or "Source processed successfully."
            ),
        }
    )
@login_required
@require_POST
def generate_ask_output(request, source_id):

    source = get_object_or_404(
        AskSource,
        id=source_id,
        user=request.user
    )

    output_type = request.POST.get(
        "output_type"
    )

    allowed = {
        "summary",
        "notes",
        "audio",
        "video",
        "quiz",
    }

    if output_type not in allowed:

        return JsonResponse(
            {
                "success": False,
                "error": "Invalid output type."
            },
            status=400
        )

    from .services.ai_service import (
        generate_from_source
    )

    source_text = (
        source.transcript
        or source.text
    )

    if not source_text.strip():

        return JsonResponse({
            "success": False,
            "error": source.access_message
            or "There is no readable source content."
        })

    try:

        content = generate_from_source(
            source_text,
            output_type
        )

        output = AskOutput.objects.create(
            source=source,
            output_type=output_type,
            content=content
        )

        return JsonResponse({
            "success": True,
            "output": {
                "id": output.id,
                "type": output.output_type,
                "content": output.content
            }
        })

    except Exception as error:

        return JsonResponse(
            {
                "success": False,
                "error": str(error)
            },
            status=500
        )
def ask_ai(request):
    if request.method == "POST":
        question = request.POST.get("question")

        client = OpenAI(api_key=settings.OPENAI_API_KEY)

        response = client.responses.create(
            model="gpt-5-mini",
            input=question
        )

        answer = response.output_text

        return JsonResponse({
            "answer": answer
        })

    return JsonResponse({
        "error": "Only POST requests are allowed."
    })

def landing(request):
    return render(
        request,
        "landing.html"
    )


def register_view(request):

    if request.method == "POST":

        form = RegisterForm(request.POST)

        if form.is_valid():

            name = form.cleaned_data["full_name"]
            email = form.cleaned_data["email"]
            password = form.cleaned_data["password"]

            otp = str(
                random.randint(
                    100000,
                    999999
                )
            )

            request.session["register_name"] = name
            request.session["register_email"] = email
            request.session["register_password"] = password
            request.session["register_otp"] = otp

            send_mail(
                "Cognify Email Verification",
                f"Your Cognify verification OTP is: {otp}",
                None,
                [email],
                fail_silently=False,
            )

            return redirect("verify")

    else:
        form = RegisterForm()

    return render(
        request,
        "registration/register.html",
        {"form": form}
    )


def verify_otp(request):

    email = request.session.get(
        "register_email"
    )

    if not email:
        return redirect("register")

    if request.method == "POST":

        entered_otp = request.POST.get(
            "otp"
        )

        saved_otp = request.session.get(
            "register_otp"
        )

        if entered_otp == saved_otp:

            name = request.session.get(
                "register_name"
            )

            password = request.session.get(
                "register_password"
            )

            user = User.objects.create_user(
                username=email,
                email=email,
                password=password,
                first_name=name
            )

            login(
                request,
                user
            )

            request.session.pop(
                "register_name",
                None
            )

            request.session.pop(
                "register_email",
                None
            )

            request.session.pop(
                "register_password",
                None
            )

            request.session.pop(
                "register_otp",
                None
            )

            return redirect("dashboard")

        messages.error(
            request,
            "Invalid OTP. Please try again."
        )

    return render(
        request,
        "registration/verify.html"
    )


def login_view(request):

    if request.method == "POST":

        form = LoginForm(request.POST)

        if form.is_valid():

            email = form.cleaned_data["email"]
            password = form.cleaned_data["password"]

            user = authenticate(
                request,
                username=email,
                password=password
            )

            if user is not None:

                login(
                    request,
                    user
                )

                return redirect(
                    "dashboard"
                )

            messages.error(
                request,
                "Invalid email or password."
            )

    else:
        form = LoginForm()

    return render(
        request,
        "registration/login.html",
        {"form": form}
    )


def logout_view(request):

    logout(request)

    return redirect("landing")


@login_required
def dashboard(request):

    notebooks = Notebook.objects.filter(
        owner=request.user
    ).order_by("-updated_at")

    results = GeneratedResult.objects.filter(
        user=request.user
    ).order_by("-created_at")[:5]

    total_sources = sum(
        notebook.sources.count()
        for notebook in notebooks
    )

    total_chats = sum(
        notebook.messages.count()
        for notebook in notebooks
    )

    total_saved = GeneratedResult.objects.filter(
        user=request.user,
        is_saved=True
    ).count()

    return render(
        request,
        "core/dashboard.html",
        {
            "notebooks": notebooks,
            "results": results,

            "total_sources": total_sources,

            "total_chats": total_chats,

            "total_saved": total_saved,
        }
    )
@login_required
def delete_notebook(request, notebook_id):

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    if request.method == "POST":

        notebook.delete()

        messages.success(
            request,
            "Notebook deleted successfully."
        )

        return redirect("dashboard")

    return redirect("dashboard")


@login_required
def my_notebooks(request):
    notebooks = Notebook.objects.filter(
        owner=request.user
    ).order_by("-updated_at")

    return render(
        request,
        "core/my_notebooks.html",
        {
            "notebooks": notebooks
        }
    )
@login_required
def notebook_detail(request, notebook_id):
    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    sources = notebook.sources.all().order_by("-uploaded_at")

    return render(
        request,
        "core/workspace.html",
        {
            "notebook": notebook,
            "sources": sources,
        }
    )
def extract_text_from_file(uploaded_file):

    file_name = uploaded_file.name.lower()

    # TXT
    if file_name.endswith(".txt"):
        return uploaded_file.read().decode("utf-8", errors="ignore")

    # PDF
    elif file_name.endswith(".pdf"):
        reader = PdfReader(uploaded_file)

        text = ""

        for page in reader.pages:
            text += page.extract_text() or ""

        return text

    # DOCX
    elif file_name.endswith(".docx"):
        document = Document(uploaded_file)

        text = []

        for paragraph in document.paragraphs:
            text.append(paragraph.text)

        return "\n".join(text)

    # PPTX
    elif file_name.endswith(".pptx"):
        presentation = Presentation(uploaded_file)

        text = []

        for slide in presentation.slides:
            for shape in slide.shapes:

                if hasattr(shape, "text"):
                    text.append(shape.text)

        return "\n".join(text)

    return ""
def create_video_from_text(text, output_path):
    """
    Create a multi-scene educational MP4.

    Each scene contains:
    - Scene number
    - Short educational text
    - Its own narration
    - H.264 video
    - AAC audio
    """

    import re
    import shutil

    output_path = Path(output_path)
    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    # -----------------------------------------
    # CLEAN INPUT
    # -----------------------------------------

    clean_text = " ".join(
        text.replace("\r", " ")
        .replace("\n", " ")
        .split()
    )

    if not clean_text:
        clean_text = (
            "Welcome to your Cognify AI learning video."
        )

    # -----------------------------------------
    # SPLIT INTO SENTENCES
    # -----------------------------------------

    sentences = re.split(
        r"(?<=[.!?])\s+",
        clean_text
    )

    sentences = [
        sentence.strip()
        for sentence in sentences
        if sentence.strip()
    ]

    # -----------------------------------------
    # BUILD SCENES
    # -----------------------------------------

    scenes = []
    current_scene = []
    current_words = 0

    for sentence in sentences:

        sentence_words = len(
            sentence.split()
        )

        if (
            current_scene
            and current_words + sentence_words > 90
        ):
            scenes.append(
                " ".join(current_scene)
            )

            current_scene = []
            current_words = 0

        current_scene.append(sentence)
        current_words += sentence_words

    if current_scene:
        scenes.append(
            " ".join(current_scene)
        )

    # Maximum 5 scenes
    scenes = scenes[:5]

    if not scenes:
        scenes = [clean_text]

    print(
        "Cognify scenes:",
        len(scenes)
    )

    # -----------------------------------------
    # TEMP DIRECTORY
    # -----------------------------------------

    temp_directory = (
        output_path.parent
        / f"video_scenes_{output_path.stem}"
    )

    temp_directory.mkdir(
        parents=True,
        exist_ok=True
    )

    scene_files = []

    try:

        # -----------------------------------------
        # WINDOWS TTS + FFMPEG
        # -----------------------------------------

        for index, scene_text in enumerate(
            scenes,
            start=1
        ):

            print(
                f"Creating scene {index}/{len(scenes)}..."
            )

            text_path = (
                temp_directory
                / f"scene_{index}.txt"
            )

            wav_path = (
                temp_directory
                / f"scene_{index}.wav"
            )

            mp4_path = (
                temp_directory
                / f"scene_{index}.mp4"
            )

            # -------------------------------------
            # SHORT ON-SCREEN VERSION
            # -------------------------------------

            scene_sentences = re.split(
                r"(?<=[.!?])\s+",
                scene_text
            )

            display_text = " ".join(
                scene_sentences[:2]
            )

            # Wrap text
            words = display_text.split()

            lines = []
            current_line = ""

            for word in words:

                test_line = (
                    current_line + " " + word
                ).strip()

                if len(test_line) <= 48:

                    current_line = test_line

                else:

                    if current_line:
                        lines.append(
                            current_line
                        )

                    current_line = word

            if current_line:
                lines.append(
                    current_line
                )

            display_text = "\n".join(
                lines[:6]
            )

            text_path.write_text(
                display_text,
                encoding="utf-8"
            )
            narration_path.write_text(
            scene_text,
            encoding="utf-8"
            )

            # -------------------------------------
            # GENERATE NARRATION
            # -------------------------------------

            powershell_script = f"""
Add-Type -AssemblyName System.Speech

$synth = New-Object System.Speech.Synthesis.SpeechSynthesizer

$synth.Rate = 0
$synth.Volume = 100

$text = Get-Content -Raw -LiteralPath '{text_path.as_posix()}'

$synth.SetOutputToWaveFile(
    '{wav_path.as_posix()}'
)

$synth.Speak(
    '{scene_text.replace("'", "''")}'
)

$synth.Dispose()
"""

            subprocess.run(
                [
                    "powershell.exe",
                    "-NoProfile",
                    "-ExecutionPolicy",
                    "Bypass",
                    "-Command",
                    powershell_script
                ],
                check=True
            )

            # -------------------------------------
            # FFmpeg TEXT PATHS
            # -------------------------------------

            font_file = (
                "C:/Windows/Fonts/arial.ttf"
            )

            font_file = (
                font_file
                .replace("\\", "/")
                .replace(":", "\\:")
            )

            text_file = (
                str(text_path)
                .replace("\\", "/")
                .replace(":", "\\:")
            )
            narration_path = (
            temp_directory
            / f"narration_{index}.txt"
            )

            # -------------------------------------
            # CREATE SCENE VIDEO
            # -------------------------------------

            video_filter = (
                "setpts=PTS-STARTPTS,"
                "drawbox="
                "x=0:y=0:w=iw:h=ih:"
                "color=0x171B4F@1:"
                "t=fill,"
                "drawbox="
                "x=0:y=0:w=iw:h=115:"
                "color=0x101536@1:"
                "t=fill,"
                "drawtext="
                f"fontfile='{font_file}':"
                f"text='COGNIFY AI  |  SCENE {index}':"
                "fontcolor=white:"
                "fontsize=30:"
                "x=55:"
                "y=35,"
                "drawtext="
                f"fontfile='{font_file}':"
                f"textfile='{text_file}':"
                "fontcolor=white:"
                "fontsize=34:"
                "line_spacing=16:"
                "x=70:"
                "y=200:"
                "box=1:"
                "boxcolor=0x080D25@0.55:"
                "boxborderw=35"
            )

            command = [
                "ffmpeg",
                "-y",

                "-f",
                "lavfi",

                "-i",
                "color="
                "c=0x171B4F:"
                "s=1280x720:"
                "r=24",

                "-i",
                str(wav_path),

                "-map",
                "0:v:0",

                "-map",
                "1:a:0",

                "-fflags",
                "+genpts",

                "-vf",
                video_filter,

                "-c:v",
                "libx264",

                "-preset",
                "ultrafast",

                "-tune",
                "stillimage",

                "-crf",
                "28",

                "-pix_fmt",
                "yuv420p",

                "-r",
                "24",

                "-fps_mode",
                "cfr",

                "-c:a",
                "aac",

                "-b:a",
                "128k",

                "-ar",
                "48000",

                "-ac",
                "2",

                "-af",
                "aresample=async=1:first_pts=0",

                "-shortest",

                "-avoid_negative_ts",
                "make_zero",

                "-video_track_timescale",
                "90000",

                "-movflags",
                "+faststart",

                "-metadata:s:v:0",
                "rotate=0",

                str(mp4_path)
            ]

            subprocess.run(
                command,
                check=True
            )

            if not mp4_path.exists():
                raise FileNotFoundError(
                    f"Scene {index} was not created."
                )

            scene_files.append(
                mp4_path
            )

        # -----------------------------------------
        # CREATE CONCAT LIST
        # -----------------------------------------

        concat_file = (
            temp_directory
            / "scenes.txt"
        )

        with open(
            concat_file,
            "w",
            encoding="utf-8"
        ) as file:

            for scene_file in scene_files:

                scene_path = (
                    str(scene_file)
                    .replace("\\", "/")
                )

                file.write(
                    f"file '{scene_path}'\n"
                )

        # -----------------------------------------
        # JOIN ALL SCENES
        # -----------------------------------------

        print(
            "Combining all Cognify scenes..."
        )

        concat_command = [
            "ffmpeg",
            "-y",

            "-f",
            "concat",

            "-safe",
            "0",

            "-i",
            str(concat_file),

            "-c",
            "copy",

            "-movflags",
            "+faststart",

            "-avoid_negative_ts",
            "make_zero",

            str(output_path)
        ]

        subprocess.run(
            concat_command,
            check=True
        )

        if not output_path.exists():
            raise FileNotFoundError(
                "Final Cognify video was not created."
            )

        if output_path.stat().st_size == 0:
            raise ValueError(
                "Final Cognify video is empty."
            )

        print(
            "MULTI-SCENE VIDEO CREATED:",
            output_path
        )

    finally:

        # -----------------------------------------
        # CLEAN TEMP FILES
        # -----------------------------------------

        if temp_directory.exists():

            shutil.rmtree(
                temp_directory,
                ignore_errors=True
            )

    return output_path
@login_required
def add_source(request, notebook_id):

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    if request.method == "POST":

        uploaded_file = request.FILES.get("file")

        if uploaded_file:

            extracted_text = extract_text_from_file(
                uploaded_file
            )

            Source.objects.create(
                notebook=notebook,
                file=uploaded_file,
                extracted_text=extracted_text
            )

        return redirect(
            "notebook_detail",
            notebook_id=notebook.id
        )

    return redirect(
        "notebook_detail",
        notebook_id=notebook.id
    )

@login_required
def create_notebook(request):

    if request.method == "POST":

        form = NotebookForm(request.POST)

        if form.is_valid():

            notebook = form.save(commit=False)

            notebook.owner = request.user

            notebook.save()

            messages.success(
                request,
                "Notebook created successfully."
            )

            return redirect(
                "my_notebooks"
            )

    else:
        form = NotebookForm()

    return render(
        request,
        "core/create_notebook.html",
        {
            "form": form
        }
    )


@login_required
def profile_settings(request):

    user = request.user

    if request.method == "POST":

        full_name = request.POST.get(
            "full_name",
            ""
        ).strip()

        email = request.POST.get(
            "email",
            ""
        ).strip()

        if not full_name:
            messages.error(
                request,
                "Full name cannot be empty."
            )

        elif not email:
            messages.error(
                request,
                "Email cannot be empty."
            )

        else:

            user.first_name = full_name
            user.email = email

            # Your project uses email as username.
            user.username = email

            user.save()

            messages.success(
                request,
                "Profile updated successfully."
            )

            return redirect(
                "profile_settings"
            )

    return render(
        request,
        "core/profile.html",
        {
            "user": user
        }
    )
@login_required
def security_settings(request):

    if request.method == "POST":

        form = PasswordChangeForm(
            request.user,
            request.POST
        )

        if form.is_valid():

            user = form.save()

            update_session_auth_hash(
                request,
                user
            )

            messages.success(
                request,
                "Password changed successfully."
            )

            return redirect(
                "security_settings"
            )

    else:

        form = PasswordChangeForm(
            request.user
        )

    return render(
        request,
        "core/security.html",
        {
            "form": form
        }
    )

@login_required
def workspace(request, notebook_id):

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    sources = notebook.sources.all().order_by("-uploaded_at")

    return render(
        request,
        "core/workspace.html",
        {
            "notebook": notebook,
            "sources": sources,
        }
    )

@login_required
def delete_notebook(request, notebook_id):

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    if request.method == "POST":

        notebook.delete()

        messages.success(
            request,
            "Notebook deleted successfully."
        )

    return redirect("my_notebooks")
@login_required
def settings_view(request):

    user = request.user

    if request.method == "POST":

        action = request.POST.get("action")

        # ==============================================
        # PROFILE UPDATE
        # ==============================================

        if action == "profile":

            full_name = request.POST.get(
                "full_name",
                ""
            ).strip()

            email = request.POST.get(
                "email",
                ""
            ).strip()

            if not full_name:

                messages.error(
                    request,
                    "Full name is required."
                )

                return redirect("settings")

            if not email:

                messages.error(
                    request,
                    "Email is required."
                )

                return redirect("settings")

            # Check duplicate email
            email_exists = User.objects.filter(
                email__iexact=email
            ).exclude(
                id=user.id
            ).exists()

            if email_exists:

                messages.error(
                    request,
                    "This email is already registered."
                )

                return redirect("settings")

            user.first_name = full_name
            user.email = email

            # Your login system uses email as username
            user.username = email

            user.save()

            messages.success(
                request,
                "Profile updated successfully."
            )

            return redirect("settings")


        # ==============================================
        # PASSWORD UPDATE
        # ==============================================

        if action == "password":

            current_password = request.POST.get(
                "current_password",
                ""
            )

            new_password = request.POST.get(
                "new_password",
                ""
            )

            confirm_password = request.POST.get(
                "confirm_password",
                ""
            )

            if not user.check_password(
                current_password
            ):

                messages.error(
                    request,
                    "Current password is incorrect."
                )

                return redirect("settings")

            if len(new_password) < 8:

                messages.error(
                    request,
                    "New password must contain at least 8 characters."
                )

                return redirect("settings")

            if new_password != confirm_password:

                messages.error(
                    request,
                    "New passwords do not match."
                )

                return redirect("settings")

            user.set_password(
                new_password
            )

            user.save()

            update_session_auth_hash(
                request,
                user
            )

            messages.success(
                request,
                "Password changed successfully."
            )

            return redirect("settings")

    return render(
        request,
        "core/settings.html",
        {
            "user": user
        }
    )
@login_required
def profile(request):

    return render(
        request,
        "core/profile.html",
        {
            "user": request.user
        }
    )
@login_required
def generate_output(request, notebook_id, kind):

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    sources = notebook.sources.all().order_by("-uploaded_at")

    text_parts = []

    for source in sources:
        if source.extracted_text:
            text_parts.append(
                source.extracted_text
            )

    if not text_parts:
        messages.error(
            request,
            "No extracted text is available for this notebook."
        )

        return redirect(
            "workspace",
            notebook_id=notebook.id
        )

    combined_text = "\n\n".join(text_parts)

    titles = {
        "faq": "Frequently Asked Questions",
        "notes": "Study Notes",
        "flashcards": "Flashcards",
        "quiz": "AI Generated Quiz",
        "diagram": "AI Diagram",
        "video": "Video Overview",
        "audio": "Audio Overview",
    }

    title = titles.get(
        kind,
        "AI Generated Result"
    )

    prompts = {

        "faq": (
            "Create 10 frequently asked questions and answers "
            "from the uploaded notebook sources. "
            "Use only the provided information. "
            "Format each item as:\n\n"
            "Q1. Question\n"
            "Answer: Answer\n\n"
        ),

        "notes": (
            "Create clear and well-organized study notes "
            "from the uploaded notebook sources. "
            "Use headings, bullet points, definitions, "
            "important concepts, and examples where available."
        ),

        "flashcards": (
    "Create 10 useful study flashcards from the uploaded "
    "notebook sources. Use only information from the sources.\n\n"
    "Format them exactly like this:\n\n"
    "Card 1\n"
    "Question: [question]\n"
    "Answer: [short answer]\n\n"
    "Card 2\n"
    "Question: [question]\n"
    "Answer: [short answer]\n\n"
    "Continue until Card 10.\n\n"
    "Keep questions clear and answers concise. "
    "Focus on important concepts, definitions, facts, "
    "processes, and examples."
),

        "quiz": (
            "Create a 10-question quiz from the uploaded "
            "notebook sources. Include the correct answer "
            "after each question."
        ),

        "diagram": (
            "Create a clear text-based concept diagram showing "
            "the important concepts and their relationships "
            "from the uploaded notebook sources."
        ),

        "video": (
            "Create a clear educational video script based "
            "only on the uploaded notebook sources."
        ),

        "audio": (
            "Create an educational audio/podcast-style script "
            "based only on the uploaded notebook sources."
        ),
    }

    prompt = prompts.get(
        kind,
        "Create a useful learning resource from the uploaded sources."
    )

    try:

        response = gemini_generate(
    system_prompt='You are Cognify AI, a learning assistant. Use only the uploaded notebook sources.',
    user_prompt=prompt + '\n\nUploaded notebook sources:\n\n' + combined_text
)

        content = content

    except Exception as e:

        messages.error(
            request,
            "AI generation failed: " + str(e)
        )

        return redirect(
            "workspace",
            notebook_id=notebook.id
        )

    result = GeneratedResult.objects.create(
        notebook=notebook,
        user=request.user,
        kind=kind,
        title=title,
        content=content
    )
    if kind == "video":

        video_filename = f"video_{result.id}.mp4"

        video_directory = (
            Path(settings.MEDIA_ROOT)
            / "generated_videos"
        )

        video_directory.mkdir(
            parents=True,
            exist_ok=True
        )

        video_path = (
            video_directory
            / video_filename
        )

        try:

            create_video_from_text(
                content,
                video_path
            )

            result.video_file.name = (
                f"generated_videos/{video_filename}"
            )

            result.save(
                update_fields=["video_file"]
            )

        except Exception as e:

            messages.error(
                request,
                "Video generation failed: " + str(e)
            )

            return redirect(
                "workspace",
                notebook_id=notebook.id
            )

        return render(
            request,
            "core/video.html",
            {
                "notebook": notebook,
                "result": result
            }
        )

    return render(
        request,
        "core/result.html",
        {
            "notebook": notebook,
            "result": result
        }
    )

@login_required
def history(request):

    results = GeneratedResult.objects.filter(
        user=request.user
    ).order_by("-created_at")

    return render(
        request,
        "core/history.html",
        {
            "results": results
        }
    )


@login_required
def saved_results(request):

    results = GeneratedResult.objects.filter(
        user=request.user,
        is_saved=True
    ).order_by("-created_at")

    return render(
        request,
        "core/saved.html",
        {
            "results": results
        }
    )


@login_required
def settings_page(request):

    # -------------------------
    # CHANGE EMAIL
    # -------------------------

    if request.method == "POST":

        email = request.POST.get(
            "email",
            ""
        ).strip()

        if not email:

            messages.error(
                request,
                "Please enter an email address."
            )

            return redirect("settings_page")


        request.user.email = email

        request.user.save()


        messages.success(
            request,
            "Email address updated successfully."
        )


        return redirect(
            "settings_page"
        )


    # -------------------------
    # NOTEBOOK COUNT
    # -------------------------

    notebooks_count = Notebook.objects.filter(
        owner=request.user
    ).count()


    # -------------------------
    # STORAGE
    # -------------------------

    total_bytes = 0

    uploaded_files = 0


    sources = Source.objects.filter(
        notebook__owner=request.user
    )


    for source in sources:

        for field in source._meta.fields:

            if isinstance(field, FileField):

                file_value = getattr(
                    source,
                    field.name,
                    None
                )


                if file_value:

                    try:

                        total_bytes += file_value.size

                        uploaded_files += 1

                    except (
                        ValueError,
                        OSError
                    ):

                        pass


    # -------------------------
    # FORMAT STORAGE
    # -------------------------

    if total_bytes < 1024:

        storage_used = (
            f"{total_bytes} B"
        )

    elif total_bytes < 1024 * 1024:

        storage_used = (
            f"{total_bytes / 1024:.2f} KB"
        )

    elif total_bytes < 1024 * 1024 * 1024:

        storage_used = (
            f"{total_bytes / (1024 * 1024):.2f} MB"
        )

    else:

        storage_used = (
            f"{total_bytes / (1024 * 1024 * 1024):.2f} GB"
        )


    # -------------------------
    # STORAGE LIMIT
    # 1 GB
    # -------------------------

    storage_limit = (
        1024 * 1024 * 1024
    )


    storage_percentage = min(
        round(
            (
                total_bytes /
                storage_limit
            ) * 100,
            1
        ),
        100
    )


    # -------------------------
    # PAGE
    # -------------------------

    return render(
        request,
        "core/settings.html",
        {
            "notebooks_count":
                notebooks_count,

            "storage_used":
                storage_used,

            "uploaded_files":
                uploaded_files,

            "storage_percentage":
                storage_percentage,
        }
    )
@login_required
def upload_source(request, notebook_id):

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    if request.method == "POST":

        uploaded_file = request.FILES.get("file")

        if not uploaded_file:
            messages.error(
                request,
                "Please select a file."
            )

            return redirect(
                "workspace",
                notebook_id=notebook.id
            )

        file_name = uploaded_file.name

        extension = os.path.splitext(
            file_name
        )[1].lower()

        allowed_extensions = [
            ".pdf",
            ".docx",
            ".pptx",
            ".txt",
        ]

        if extension not in allowed_extensions:

            messages.error(
                request,
                "Only PDF, DOCX, PPTX and TXT files are supported."
            )

            return redirect(
                "workspace",
                notebook_id=notebook.id
            )

        extracted_text = ""

        try:

            # PDF
            if extension == ".pdf":

                reader = PdfReader(
                    uploaded_file
                )

                for page in reader.pages:

                    text = page.extract_text()

                    if text:
                        extracted_text += text + "\n"

            # DOCX
            elif extension == ".docx":

                document = Document(
                    uploaded_file
                )

                for paragraph in document.paragraphs:

                    extracted_text += (
                        paragraph.text + "\n"
                    )

            # PPTX
            elif extension == ".pptx":

                presentation = Presentation(
                    uploaded_file
                )

                for slide in presentation.slides:

                    for shape in slide.shapes:

                        if hasattr(
                            shape,
                            "text"
                        ):

                            extracted_text += (
                                shape.text + "\n"
                            )

            # TXT
            elif extension == ".txt":

                extracted_text = (
                    uploaded_file
                    .read()
                    .decode(
                        "utf-8",
                        errors="ignore"
                    )
                )

            source = Source.objects.create(
                notebook=notebook,
                file=uploaded_file,
                extracted_text=extracted_text
            )

            messages.success(
                request,
                f"{file_name} uploaded successfully."
            )

        except Exception as e:

            messages.error(
                request,
                f"Could not process the file: {str(e)}"
            )

        return redirect(
            "workspace",
            notebook_id=notebook.id
        )

    return redirect(
        "workspace",
        notebook_id=notebook.id
    )
def create_local_summary(text, max_sentences=8):
    import re

    sentences = re.split(r'(?<=[.!?])\s+', text.strip())

    sentences = [
        sentence.strip()
        for sentence in sentences
        if len(sentence.strip()) > 30
    ]

    if not sentences:
        return text[:1000]

    keywords = [
        "important",
        "definition",
        "defined",
        "because",
        "therefore",
        "example",
        "method",
        "process",
        "result",
        "concept",
        "principle",
        "advantage",
        "disadvantage",
    ]

    scored_sentences = []

    for index, sentence in enumerate(sentences):
        lower_sentence = sentence.lower()

        score = sum(
            1 for keyword in keywords
            if keyword in lower_sentence
        )

        if index < 3:
            score += 1

        scored_sentences.append(
            (score, index, sentence)
        )

    selected = sorted(
        scored_sentences,
        key=lambda item: (-item[0], item[1])
    )[:max_sentences]

    selected.sort(key=lambda item: item[1])

    return " ".join(
        sentence
        for _, _, sentence in selected
    )
@login_required
def delete_source(request, source_id):

    source = get_object_or_404(
        Source,
        id=source_id,
        notebook__owner=request.user
    )

    notebook_id = source.notebook.id

    if request.method == "POST":

        if source.file:
            source.file.delete(save=False)

        source.delete()

        messages.success(
            request,
            "Source deleted successfully."
        )

    return redirect(
        "workspace",
        notebook_id=notebook_id
    )
@login_required
def generate_summary(request, notebook_id):

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    sources = notebook.sources.all().order_by("-uploaded_at")

    text_parts = []
    MAX_SOURCE_CHARS = 40000
    total_chars = 0

    for source in sources:

        if not source.extracted_text:
            continue

        remaining = MAX_SOURCE_CHARS - total_chars

        if remaining <= 0:
            break

        text = source.extracted_text[:remaining]

        text_parts.append(text)
        total_chars += len(text)

    if not text_parts:

        messages.error(
            request,
            "No extracted text is available for this notebook."
        )

        return redirect(
            "workspace",
            notebook_id=notebook.id
        )

    combined_text = "\n\n".join(text_parts)

    # Create result immediately
    result = GeneratedResult.objects.create(
        notebook=notebook,
        user=request.user,
        kind="summary",
        title="AI Summary",
        content="GENERATING..."
    )

    # Run AI generation separately
    thread = threading.Thread(
        target=generate_summary_background,
        args=(result.id, combined_text),
        daemon=True
    )

    thread.start()

    # OPEN PAGE IMMEDIATELY
    return render(
        request,
        "core/result.html",
        {
            "notebook": notebook,
            "result": result,
            "generating": True
        }
    )
def generate_summary_background(result_id, combined_text):

    try:

        # Keep the AI input small for fast generation
        combined_text = combined_text[:6000]

        response = gemini_generate(
    system_prompt='You are Cognify AI. Summarize the study material in concise bullet points. Include only the most important concepts, definitions and facts. Do not repeat information.',
    user_prompt=combined_text
)

        summary = content

        GeneratedResult.objects.filter(
            id=result_id
        ).update(
            content=summary
        )

    except Exception as e:

        GeneratedResult.objects.filter(
            id=result_id
        ).update(
            content="ERROR: " + str(e)
        )
@login_required
def summary_status(request, result_id):

    result = get_object_or_404(
        GeneratedResult,
        id=result_id,
        user=request.user
    )

    return JsonResponse({
        "content": result.content,
        "ready": result.content != "GENERATING..."
    })
@login_required
def result_detail(request, result_id):
    result = get_object_or_404(
        GeneratedResult,
        id=result_id,
        notebook__owner=request.user
    )

    return render(
        request,
        "core/result_detail.html",
        {
            "result": result
        }
    )
@login_required
def save_result(request, result_id):

    result = get_object_or_404(
        GeneratedResult,
        id=result_id,
        user=request.user
    )

    if request.method == "POST":

        result.is_saved = True
        result.save()

        messages.success(
            request,
            "Result saved successfully."
        )

    return redirect(
        "workspace",
        notebook_id=result.notebook.id
    )
@login_required
def ask_question(request, notebook_id):

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    sources = notebook.sources.all().order_by("-uploaded_at")

    text_parts = []

    for source in sources:
        if source.extracted_text:
            text_parts.append(
                source.extracted_text
            )

    if request.method == "POST":

        question = request.POST.get(
            "question",
            ""
        ).strip()

        if not question:
            messages.error(
                request,
                "Please enter a question."
            )

            return redirect(
                "ask_question",
                notebook_id=notebook.id
            )

        if not text_parts:
            messages.error(
                request,
                "No extracted text is available "
                "for this notebook."
            )

            return redirect(
                "workspace",
                notebook_id=notebook.id
            )

        combined_text = "\n\n".join(
            text_parts
        )

        try:

            response = gemini_generate(
    system_prompt="You are Cognify AI, a learning assistant. Answer the user's question using only the uploaded notebook sources. If the answer is not found in the sources, clearly say that it was not found.",
    user_prompt=f'\nUploaded notebook sources:\n\n{combined_text}\n\nUser question:\n\n{question}\n'
)

            answer = content

        except Exception as e:

            answer = (
                "Cognify AI encountered an error: "
                + str(e)
            )

        return render(request, "core/workspace.html", {
    "notebook": notebook,
    "sources": sources,
    "question": question,
    "answer": answer,
})
    return render(
        request,
        "core/ask_question.html",
        {
            "notebook": notebook,
        }
    )
@login_required
def generate_faq(request, notebook_id):

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    sources = notebook.sources.all().order_by("-uploaded_at")

    text_parts = []

    for source in sources:
        if source.extracted_text:
            text_parts.append(source.extracted_text)

    if not text_parts:
        messages.error(
            request,
            "No extracted text is available for this notebook."
        )

        return redirect(
            "workspace",
            notebook_id=notebook.id
        )

    combined_text = "\n\n".join(text_parts)

    try:

        response = gemini_generate(
    system_prompt='You are Cognify AI. Create exactly 10 useful FAQs from the provided notebook sources. Use only information from the sources.',
    user_prompt='Create exactly 10 FAQs.\n\nUse this exact format:\n\nQ1: Question\nA1: Answer\n\nQ2: Question\nA2: Answer\n\nContinue until Q10.\n\nNotebook sources:\n\n' + combined_text
)

        content = content

    except Exception as e:

        messages.error(
            request,
            "FAQ generation failed: " + str(e)
        )

        return redirect(
            "workspace",
            notebook_id=notebook.id
        )


    result = GeneratedResult.objects.create(
        notebook=notebook,
        user=request.user,
        kind="faq",
        title="Frequently Asked Questions",
        content=content
    )


    import re

    faqs = []

    pattern = re.compile(
        r"Q\d+\s*:\s*(.*?)\s*"
        r"A\d+\s*:\s*(.*?)(?="
        r"\s*Q\d+\s*:|$)",
        re.IGNORECASE | re.DOTALL
    )

    matches = pattern.findall(content)

    for question, answer in matches:

        question = question.strip()
        answer = answer.strip()

        if question and answer:

            faqs.append({
                "question": question,
                "answer": answer
            })


    return render(
        request,
        "core/faq.html",
        {
            "notebook": notebook,
            "result": result,
            "faqs": faqs,
        }
    )
@login_required
def generate_notes(request, notebook_id):

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    sources = notebook.sources.all().order_by("-uploaded_at")

    text_parts = []

    for source in sources:
        if source.extracted_text:
            text_parts.append(source.extracted_text)

    if not text_parts:
        messages.error(
            request,
            "No extracted text is available for this notebook."
        )

        return redirect(
            "workspace",
            notebook_id=notebook.id
        )

    combined_text = "\n\n".join(text_parts)

    try:

        response = gemini_generate(
    system_prompt='You are Cognify AI, a study assistant. Create clear, organized study notes using only the supplied notebook sources.',
    user_prompt='Create detailed but easy-to-understand study notes.\n\nUse these sections:\nIntroduction\nKey Points\nImportant Definitions\nExamples\nApplications\nSummary\n\nUse headings and bullet points. Focus on information useful for exam preparation.\n\nNotebook sources:\n\n' + combined_text
)

        content = content

    except Exception as e:

        messages.error(
            request,
            "Study Notes generation failed: " + str(e)
        )

        return redirect(
            "workspace",
            notebook_id=notebook.id
        )


    result = GeneratedResult.objects.create(
        notebook=notebook,
        user=request.user,
        kind="notes",
        title="Study Notes",
        content=content
    )


    return render(
        request,
        "core/study_notes.html",
        {
            "notebook": notebook,
            "result": result,
        }
    )
@login_required
def generate_flashcards(request, notebook_id):

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    sources = notebook.sources.all().order_by("-uploaded_at")

    text_parts = []

    for source in sources:

        if source.extracted_text:
            text_parts.append(
                source.extracted_text
            )

    if not text_parts:

        messages.error(
            request,
            "No extracted text is available for this notebook."
        )

        return redirect(
            "workspace",
            notebook_id=notebook.id
        )

    combined_text = "\n\n".join(text_parts)

    try:

        response = gemini_generate(
    system_prompt='You are Cognify AI, a study assistant. Create exactly 10 useful flashcards from the supplied notebook sources.',
    user_prompt='Create exactly 10 flashcards.\n\nReturn ONLY valid JSON in this format:\n\n[{"question":"Question 1","answer":"Answer 1"},{"question":"Question 2","answer":"Answer 2"}]\n\nDo not include markdown.\nDo not include explanations outside JSON.\n\nNotebook sources:\n\n' + combined_text
)

        content = content.strip()

        import json

        if content.startswith("```"):

            content = content.replace(
                "```json",
                ""
            ).replace(
                "```",
                ""
            ).strip()

        flashcards = json.loads(content)

    except Exception as e:

        messages.error(
            request,
            "Flashcard generation failed: " + str(e)
        )

        return redirect(
            "workspace",
            notebook_id=notebook.id
        )


    result = GeneratedResult.objects.create(
        notebook=notebook,
        user=request.user,
        kind="flashcards",
        title="Flashcards",
        content=content
    )


    return render(
        request,
        "core/flashcards.html",
        {
            "notebook": notebook,
            "result": result,
            "flashcards": flashcards,
        }
    ),
@login_required
def generate_quiz(request, notebook_id):

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    sources = notebook.sources.all().order_by("-uploaded_at")

    text_parts = []

    for source in sources:
        if source.extracted_text:
            text_parts.append(source.extracted_text)

    if not text_parts:
        messages.error(
            request,
            "No extracted text is available for this notebook."
        )
        return redirect(
            "workspace",
            notebook_id=notebook.id
        )

    combined_text = "\n\n".join(text_parts)

    try:
        response = gemini_generate(
    system_prompt='You are Cognify AI, a quiz generator. Create exactly 10 multiple choice questions using ONLY the supplied notebook content. Return ONLY valid JSON. Never use markdown. Never add text before or after the JSON.',
    user_prompt=f'\nCreate exactly 10 multiple choice questions.\n\nEach question must contain:\n- question\n- exactly 4 options\n- correct\n\nThe correct value must be a number from 0 to 3.\n\nUse exactly this JSON structure:\n\n[\n  {{\n    "question": "What is Machine Learning?",\n    "options": [\n      "Option A",\n      "Option B",\n      "Option C",\n      "Option D"\n    ],\n    "correct": 1\n  }}\n]\n\nReturn ONLY the JSON array.\n\nNotebook content:\n\n{combined_text}\n'
)

        content = content.strip()

        import json
        import re

        # Remove markdown code fences if Gemini adds them
        content = re.sub(
            r"^```(?:json)?\s*",
            "",
            content,
            flags=re.IGNORECASE
        )

        content = re.sub(
            r"\s*```$",
            "",
            content
        )

        content = content.strip()

        # Find the JSON array if Gemini added extra text
        start = content.find("[")
        end = content.rfind("]")

        if start == -1 or end == -1:
            raise ValueError(
                "Gemini did not return a valid JSON array."
            )

        content = content[start:end + 1]

        quiz_questions = json.loads(content)

        # Validate the generated structure
        if not isinstance(quiz_questions, list):
            raise ValueError(
                "Quiz data is not a list."
            )

        if len(quiz_questions) == 0:
            raise ValueError(
                "Gemini returned no quiz questions."
            )

        for question in quiz_questions:

            if "question" not in question:
                raise ValueError(
                    "A quiz question is missing the question field."
                )

            if "options" not in question:
                raise ValueError(
                    "A quiz question is missing options."
                )

            if len(question["options"]) != 4:
                raise ValueError(
                    "Each quiz question must have exactly 4 options."
                )

            if "correct" not in question:
                raise ValueError(
                    "A quiz question is missing the correct answer."
                )

            if question["correct"] not in [0, 1, 2, 3]:
                raise ValueError(
                    "Correct answer must be between 0 and 3."
                )

    except Exception as e:

        messages.error(
            request,
            f"Quiz generation failed: {e}"
        )

        return redirect(
            "workspace",
            notebook_id=notebook.id
        )

    result = GeneratedResult.objects.create(
        notebook=notebook,
        user=request.user,
        kind="quiz",
        title="AI Generated Quiz",
        content=content
    )

    return render(
        request,
        "core/quiz.html",
        {
            "notebook": notebook,
            "result": result,
            "quiz_questions": quiz_questions,
        }
    )
@login_required
def generate_diagram(request, notebook_id):

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    sources = notebook.sources.all().order_by("-uploaded_at")

    # Get the topic entered by the user
    topic = request.POST.get("topic", "").strip()

    # If no topic was submitted, show the diagram page
    if request.method != "POST" or not topic:

        return render(
            request,
            "core/diagram.html",
            {
                "notebook": notebook,
                "topic": "",
            }
        )

    # Collect uploaded source text
    text_parts = []

    for source in sources:

        if source.extracted_text:

            text_parts.append(
                source.extracted_text
            )

    if not text_parts:

        messages.error(
            request,
            "Please upload at least one source before generating a diagram."
        )

        return render(
            request,
            "core/diagram.html",
            {
                "notebook": notebook,
                "topic": topic,
            }
        )

    # Combine source text
    combined_text = "\n\n".join(text_parts)

    # Prevent extremely large prompts
    combined_text = combined_text[:12000]

    try:

        response = gemini_generate(
    system_prompt="\nYou are Cognify AI, an educational diagram generator.\n\nThe USER'S TOPIC is the most important instruction.\n\nCreate a visual learning flow specifically about the user's topic.\n\nThe uploaded notebook sources are supporting material only.\nUse them to make the diagram accurate, but DO NOT create a\ngeneric diagram about the whole notebook.\n\nThe diagram must directly explain the requested topic.\n\nFor example:\n\nTopic: Machine Learning Process\n\nPossible steps:\nData → Preprocessing → Training → Model → Prediction\n\nTopic: Photosynthesis\n\nPossible steps:\nSunlight → Chlorophyll → Water + CO2 → Glucose → Oxygen\n\nTopic: Software Development Life Cycle\n\nPossible steps:\nPlanning → Analysis → Design → Development → Testing → Deployment\n\nReturn ONLY valid JSON.\nDo not return markdown.\nDo not use ```json.\n",
    user_prompt=f'\nUSER TOPIC:\n{topic}\n\nCreate a diagram specifically for this topic.\n\nIMPORTANT:\nEvery step in the diagram must be directly related to:\n\n{topic}\n\nUse the uploaded sources only to improve factual accuracy.\n\nCreate 4 to 7 logical steps.\n\nEach step must have:\n- title\n- short description\n\nReturn exactly this format:\n\n[\n  {{\n    "title": "Step 1",\n    "description": "Short explanation."\n  }},\n  {{\n    "title": "Step 2",\n    "description": "Short explanation."\n  }}\n]\n\nUPLOADED SOURCE MATERIAL:\n\n{combined_text}\n'
)

        content = content.strip()

        import json
        import re

        # Remove markdown fences if Gemini adds them
        content = re.sub(
            r"^```(?:json)?\s*",
            "",
            content,
            flags=re.IGNORECASE
        )

        content = re.sub(
            r"\s*```$",
            "",
            content
        )

        content = content.strip()

        # Find JSON array
        start = content.find("[")
        end = content.rfind("]")

        if start == -1 or end == -1:

            raise ValueError(
                "AI did not return valid diagram data."
            )

        json_content = content[
            start:end + 1
        ]

        diagram_nodes = json.loads(
            json_content
        )

        if not isinstance(
            diagram_nodes,
            list
        ):

            raise ValueError(
                "Diagram data is not a list."
            )

        if len(diagram_nodes) < 2:

            raise ValueError(
                "AI generated too few diagram steps."
            )

        # Validate every node
        for node in diagram_nodes:

            if "title" not in node:

                raise ValueError(
                    "A diagram step is missing its title."
                )

            if "description" not in node:

                node["description"] = ""

        # Save generated result
        result = GeneratedResult.objects.create(
            notebook=notebook,
            user=request.user,
            kind="diagram",
            title=f"AI Diagram - {topic}",
            content=json_content
        )

        return render(
            request,
            "core/diagram.html",
            {
                "notebook": notebook,
                "result": result,
                "diagram_nodes": diagram_nodes,
                "topic": topic,
            }
        )

    except Exception as e:

        messages.error(
            request,
            f"Diagram generation failed: {e}"
        )

        return render(
            request,
            "core/diagram.html",
            {
                "notebook": notebook,
                "topic": topic,
            }
        )
@login_required
def save_result(request, result_id):

    result = get_object_or_404(
        GeneratedResult,
        id=result_id,
        user=request.user
    )

    if request.method == "POST":

        result.is_saved = True
        result.save()

        messages.success(
            request,
            "Result saved successfully."
        )

    return redirect(
        request.META.get(
            "HTTP_REFERER",
            "dashboard"
        )
    )
@login_required
def generate_audio(request, notebook_id):

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    sources = notebook.sources.all().order_by("-uploaded_at")

    text_parts = []

    for source in sources:
        if source.extracted_text:
            text_parts.append(source.extracted_text)

    if not text_parts:

        messages.error(
            request,
            "No extracted text is available for this notebook."
        )

        return redirect(
            "workspace",
            notebook_id=notebook.id
        )

    combined_text = "\n\n".join(text_parts)

    try:

        response = gemini_generate(
    system_prompt='You are Cognify AI. Create an audio overview script based only on the provided notebook content. Explain the topic clearly for a student. Include an introduction, main concepts, important points, and applications.',
    user_prompt='Create an audio overview for this notebook:\n\n' + combined_text
)

        audio_script = content

    except Exception as e:

        messages.error(
            request,
            f"Audio overview generation failed: {e}"
        )

        return redirect(
            "workspace",
            notebook_id=notebook.id
        )

    result = GeneratedResult.objects.create(
        notebook=notebook,
        user=request.user,
        kind="audio",
        title="Audio Overview",
        content=audio_script
    )

    return render(
        request,
        "core/audio.html",
        {
            "notebook": notebook,
            "result": result,
        }
    )


def _generate_video_background(result_id):
    close_old_connections()

    try:
        result = GeneratedResult.objects.select_related(
            "notebook"
        ).get(
            id=result_id,
            kind="video"
        )

        notebook = result.notebook

        print("====================================")
        print("COGNIFY VIDEO GENERATION STARTED")
        print("Result ID:", result.id)
        print("Notebook:", notebook.id)
        print("====================================")

        # -----------------------------------------
        # STEP 1 — Prepare learning material
        # -----------------------------------------

        result.content = "__VIDEO_STEP_1__"
        result.save(update_fields=["content"])

        sources = notebook.sources.all()

        text_parts = []

        for source in sources:
            if source.extracted_text:
                text_parts.append(source.extracted_text)

        if not text_parts:
            raise ValueError(
                "No extracted text was found in the notebook."
            )

        combined_text = "\n\n".join(text_parts)

        # Limit input so Gemini responds faster
        combined_text = combined_text[:12000]

        print("Learning material prepared.")

        # -----------------------------------------
        # STEP 2 — Generate AI narration
        # -----------------------------------------

        result.content = "__VIDEO_STEP_2__"
        result.save(update_fields=["content"])

        print("Generating video script with Gemini...")

        response = gemini_generate(
    system_prompt='You are Cognify AI, an educational learning assistant. Create a short study-video narration from the supplied material. Explain only the most important concepts clearly. Do not use markdown headings. Do not include stage directions. Keep the narration between 350 and 600 words.',
    user_prompt='Create a study video narration from the following notebook content:\n\n' + combined_text
)

        video_script = content.strip()

        if not video_script:
            raise ValueError(
                "Gemini returned an empty video script."
            )

        print(
            "Video script generated:",
            len(video_script),
            "characters"
        )

        # -----------------------------------------
        # STEP 3 — Create MP4
        # -----------------------------------------

        result.content = "__VIDEO_STEP_3__"
        result.save(update_fields=["content"])

        video_directory = (
            Path(settings.MEDIA_ROOT)
            / "generated_videos"
        )

        video_directory.mkdir(
            parents=True,
            exist_ok=True
        )

        video_filename = (
            f"cognify_video_{result.id}.mp4"
        )

        video_path = (
            video_directory / video_filename
        )

        print("Creating MP4 with FFmpeg...")

        create_video_from_text(
            video_script,
            video_path
        )

        if not video_path.exists():
            raise FileNotFoundError(
                "FFmpeg finished but the MP4 file "
                "was not created."
            )

        if video_path.stat().st_size == 0:
            raise ValueError(
                "Generated MP4 is empty."
            )

        print(
            "VIDEO FILE CREATED:",
            video_path
        )

        # -----------------------------------------
        # COMPLETE
        # -----------------------------------------

        result.video_file.name = (
            f"generated_videos/{video_filename}"
        )

        result.content = video_script

        result.save(
            update_fields=[
                "video_file",
                "content"
            ]
        )

        print("====================================")
        print("COGNIFY VIDEO GENERATION COMPLETE")
        print(
            "URL:",
            result.video_file.url
        )
        print("====================================")

    except Exception as e:

        print("====================================")
        print("COGNIFY VIDEO GENERATION FAILED")
        print(str(e))
        print("====================================")

        try:
            result.content = (
                "__VIDEO_ERROR__"
                + str(e)
            )

            result.save(
                update_fields=["content"]
            )

        except Exception:
            pass

    finally:
        close_old_connections()


@login_required
def generate_video(request, notebook_id):
    """
    Only opens the video page.
    It does NOT generate the video.
    """

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    result = GeneratedResult.objects.filter(
        notebook=notebook,
        user=request.user,
        kind="video"
    ).order_by("-created_at").first()

    return render(
        request,
        "core/video.html",
        {
            "notebook": notebook,
            "result": result,
        }
    )


@login_required
def start_video_generation(request, notebook_id):
    """
    Creates a video result and starts generation
    in a background thread.
    """

    if request.method != "POST":
        return JsonResponse(
            {
                "success": False,
                "error": "POST request required."
            },
            status=405
        )

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    # Check whether a video is already being generated
    existing = GeneratedResult.objects.filter(
        notebook=notebook,
        user=request.user,
        kind="video"
    ).order_by("-created_at").first()

    if existing:
        if (
            not existing.video_file
            and existing.content == "__VIDEO_GENERATING__"
        ):
            return JsonResponse({
                "success": True,
                "status": "generating",
                "result_id": existing.id,
            })

    # Create new result
    result = GeneratedResult.objects.create(
        notebook=notebook,
        user=request.user,
        kind="video",
        title="Video Overview",
        content="__VIDEO_GENERATING__"
    )

    # Start background generation
    thread = threading.Thread(
        target=_generate_video_background,
        args=(result.id,),
        daemon=True
    )

    thread.start()

    print(
        "Background video thread started:",
        result.id
    )

    return JsonResponse({
        "success": True,
        "status": "generating",
        "result_id": result.id,
    })


@login_required
def video_status(request, notebook_id, result_id):

    notebook = get_object_or_404(
        Notebook,
        id=notebook_id,
        owner=request.user
    )

    result = get_object_or_404(
        GeneratedResult,
        id=result_id,
        notebook=notebook,
        user=request.user,
        kind="video"
    )

    # -----------------------------------------
    # VIDEO READY
    # -----------------------------------------

    if result.video_file:

        return JsonResponse({
            "status": "ready",
            "step": 4,
            "message": "Video ready!",
            "video_url": result.video_file.url,
            "script": result.content,
        })

    # -----------------------------------------
    # ERROR
    # -----------------------------------------

    if result.content.startswith(
        "__VIDEO_ERROR__"
    ):

        return JsonResponse({
            "status": "error",
            "step": 0,
            "error": result.content.replace(
                "__VIDEO_ERROR__",
                "",
                1
            )
        })

    # -----------------------------------------
    # STEP 1
    # -----------------------------------------

    if result.content == "__VIDEO_STEP_1__":

        return JsonResponse({
            "status": "generating",
            "step": 1,
            "message": "Preparing your learning material..."
        })

    # -----------------------------------------
    # STEP 2
    # -----------------------------------------

    if result.content == "__VIDEO_STEP_2__":

        return JsonResponse({
            "status": "generating",
            "step": 2,
            "message": "Cognify AI is creating the narration..."
        })

    # -----------------------------------------
    # STEP 3
    # -----------------------------------------

    if result.content == "__VIDEO_STEP_3__":

        return JsonResponse({
            "status": "generating",
            "step": 3,
            "message": "Creating your MP4 video..."
        })

    # -----------------------------------------
    # DEFAULT
    # -----------------------------------------

    return JsonResponse({
        "status": "generating",
        "step": 1,
        "message": "Starting video generation..."
    })
@login_required
def history(request):

    query = request.GET.get("q", "").strip()
    result_type = request.GET.get("type", "").strip()

    results = GeneratedResult.objects.filter(
        user=request.user
    ).select_related(
        "notebook"
    ).order_by(
        "-created_at"
    )

    # Search
    if query:
        results = results.filter(
            Q(title__icontains=query)
            |
            Q(content__icontains=query)
            |
            Q(notebook__name__icontains=query)
        )

    # Filter by result type
    if result_type:
        results = results.filter(
            kind=result_type
        )

    return render(
        request,
        "core/history.html",
        {
            "results": results,
            "query": query,
            "result_type": result_type,
        }
    )
@login_required
def saved_results(request):

    query = request.GET.get("q", "").strip()

    results = GeneratedResult.objects.filter(
        user=request.user,
        is_saved=True
    ).select_related(
        "notebook"
    ).order_by(
        "-created_at"
    )

    if query:
        results = results.filter(
            Q(title__icontains=query)
            |
            Q(content__icontains=query)
            |
            Q(notebook__name__icontains=query)
        )

    return render(
        request,
        "core/saved.html",
        {
            "results": results,
            "query": query,
        }
    )
@login_required
def toggle_saved(request, result_id):

    result = get_object_or_404(
        GeneratedResult,
        id=result_id,
        user=request.user
    )

    result.is_saved = not result.is_saved
    result.save(update_fields=["is_saved"])

    next_url = request.POST.get("next")

    if next_url:
        return redirect(next_url)

    return redirect("history")
@login_required
def profile(request):

    profile, created = UserProfile.objects.get_or_create(
        user=request.user
    )

    if request.method == "POST":

        full_name = request.POST.get(
            "full_name",
            ""
        ).strip()

        email = request.POST.get(
            "email",
            ""
        ).strip()

        bio = request.POST.get(
            "bio",
            ""
        ).strip()

        if not full_name:
            messages.error(
                request,
                "Full name cannot be empty."
            )

            return redirect("profile")

        if not email:
            messages.error(
                request,
                "Email cannot be empty."
            )

            return redirect("profile")

        # Check whether another user already uses email
        email_exists = User.objects.filter(
            email__iexact=email
        ).exclude(
            id=request.user.id
        ).exists()

        if email_exists:

            messages.error(
                request,
                "This email is already being used."
            )

            return redirect("profile")

        request.user.first_name = full_name
        request.user.email = email
        request.user.save(
            update_fields=[
                "first_name",
                "email"
            ]
        )

        profile.bio = bio
        profile.save()

        messages.success(
            request,
            "Profile updated successfully."
        )

        return redirect("profile")

    return render(
        request,
        "core/profile.html",
        {
            "profile": profile
        }
    )
@login_required
def security(request):

    if request.method == "POST":

        current_password = request.POST.get("current_password", "")
        new_password = request.POST.get("new_password", "")
        confirm_password = request.POST.get("confirm_password", "")

        # Check current password
        if not request.user.check_password(current_password):
            messages.error(
                request,
                "Current password is incorrect."
            )

            return render(
                request,
                "core/security.html"
            )

        # Check new passwords match
        if new_password != confirm_password:
            messages.error(
                request,
                "New passwords do not match."
            )

            return render(
                request,
                "core/security.html"
            )

        # Check empty password
        if not new_password:
            messages.error(
                request,
                "Please enter a new password."
            )

            return render(
                request,
                "core/security.html"
            )

        # Validate password using Django's validators
        try:
            validate_password(
                new_password,
                request.user
            )

        except ValidationError as error:

            for message in error.messages:
                messages.error(request, message)

            return render(
                request,
                "core/security.html"
            )

        # Change password
        request.user.set_password(new_password)
        request.user.save()

        # Keep the user logged in
        update_session_auth_hash(
            request,
            request.user
        )

        messages.success(
            request,
            "Your password has been changed successfully."
        )

        return redirect("security")

    return render(
        request,
        "core/security.html"
    )
@login_required
def settings_page(request):

    user = request.user

    # =====================================================
    # UPDATE EMAIL
    # =====================================================

    if request.method == "POST":

        new_email = request.POST.get("email", "").strip()

        if not new_email:

            messages.error(
                request,
                "Please enter an email address."
            )

        elif User.objects.filter(
            email=new_email
        ).exclude(
            id=user.id
        ).exists():

            messages.error(
                request,
                "That email address is already in use."
            )

        else:

            user.email = new_email
            user.username = new_email

            user.save(
                update_fields=[
                    "email",
                    "username"
                ]
            )

            messages.success(
                request,
                "Email address updated successfully."
            )

            return redirect("settings_page")


    # =====================================================
    # GET USER'S UPLOADED FILES
    # =====================================================

    sources = Source.objects.filter(
        notebook__owner=user
    )


    # =====================================================
    # STORAGE LIMIT
    # =====================================================

    one_gb = 1024 * 1024 * 1024


    # =====================================================
    # STORAGE VARIABLES
    # =====================================================

    storage_bytes = 0

    uploaded_files_count = sources.count()


    # Category sizes

    images_bytes = 0
    videos_bytes = 0
    audio_bytes = 0
    documents_bytes = 0
    other_bytes = 0


    # =====================================================
    # CALCULATE FILE SIZES
    # =====================================================

    for source in sources:

        try:

            if not source.file:
                continue

            file_size = source.file.size

            storage_bytes += file_size


            # Get filename

            filename = source.file.name.lower()


            # IMAGE

            if filename.endswith(
                (
                    ".jpg",
                    ".jpeg",
                    ".png",
                    ".gif",
                    ".webp",
                    ".bmp",
                    ".svg"
                )
            ):

                images_bytes += file_size


            # VIDEO

            elif filename.endswith(
                (
                    ".mp4",
                    ".avi",
                    ".mkv",
                    ".mov",
                    ".webm"
                )
            ):

                videos_bytes += file_size


            # AUDIO

            elif filename.endswith(
                (
                    ".mp3",
                    ".wav",
                    ".aac",
                    ".ogg",
                    ".m4a"
                )
            ):

                audio_bytes += file_size


            # DOCUMENTS

            elif filename.endswith(
                (
                    ".pdf",
                    ".doc",
                    ".docx",
                    ".txt",
                    ".ppt",
                    ".pptx",
                    ".xls",
                    ".xlsx",
                    ".csv"
                )
            ):

                documents_bytes += file_size


            # OTHER

            else:

                other_bytes += file_size


        except (ValueError, OSError):

            pass


    # =====================================================
    # FORMAT STORAGE SIZE
    # =====================================================

    def format_size(size):

        if size == 0:
            return "0 MB"

        if size < 1024:

            return f"{size} B"

        if size < 1024 * 1024:

            return f"{size / 1024:.1f} KB"

        if size < 1024 * 1024 * 1024:

            return f"{size / (1024 * 1024):.1f} MB"

        return f"{size / (1024 * 1024 * 1024):.2f} GB"


    # =====================================================
    # TOTAL STORAGE
    # =====================================================

    storage_used = format_size(
        storage_bytes
    )


    # =====================================================
    # STORAGE PERCENTAGE
    # =====================================================

    storage_percentage = (
        storage_bytes / one_gb
    ) * 100

    storage_percentage = min(
        storage_percentage,
        100
    )

    storage_percentage = round(
        storage_percentage,
        1
    )


    # =====================================================
    # FREE STORAGE
    # =====================================================

    free_bytes = max(
        one_gb - storage_bytes,
        0
    )

    free_storage = format_size(
        free_bytes
    )


    # =====================================================
    # CATEGORY STORAGE
    # =====================================================

    images_storage = format_size(
        images_bytes
    )

    videos_storage = format_size(
        videos_bytes
    )

    audio_storage = format_size(
        audio_bytes
    )

    documents_storage = format_size(
        documents_bytes
    )

    other_storage = format_size(
        other_bytes
    )


    # =====================================================
    # NOTEBOOK COUNT
    # =====================================================

    notebooks_count = Notebook.objects.filter(
        owner=user
    ).count()


    # =====================================================
    # SEND DATA TO SETTINGS.HTML
    # =====================================================

    return render(
        request,
        "core/settings.html",
        {

            # Total storage
            "storage_used": storage_used,

            "storage_percentage":
                storage_percentage,

            "free_storage":
                free_storage,

            # Files
            "uploaded_files":
                uploaded_files_count,

            # Notebooks
            "notebooks_count":
                notebooks_count,

            # Categories
            "images_storage":
                images_storage,

            "videos_storage":
                videos_storage,

            "audio_storage":
                audio_storage,

            "documents_storage":
                documents_storage,

            "other_storage":
                other_storage,
        }
    )